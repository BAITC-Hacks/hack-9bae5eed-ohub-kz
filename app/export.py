
from docx import Document
from fpdf import FPDF
from datetime import datetime

def export_docx(protocol, tasks, summary, path="protocol.docx"):
    doc = Document()
    doc.add_heading('Protokol', 0)
    doc.add_paragraph(f"Date: {datetime.now()}")
    doc.add_heading('Summary',1)
    doc.add_paragraph(summary)
    doc.add_heading('Tasks',1)
    table = doc.add_table(rows=1, cols=4)
    table.style='Light Grid'
    hdr=table.rows[0].cells
    hdr[0].text='Assignee';hdr[1].text='Task';hdr[2].text='Deadline';hdr[3].text='Source'
    for t in tasks:
        row=table.add_row().cells
        row[0].text=t.get('assignee','');row[1].text=t.get('task','');row[2].text=t.get('deadline','');row[3].text=t.get('source_quote','')[:100]
    doc.save(path)
    return path

def export_pdf(protocol, tasks, summary, path="protocol.pdf"):
    pdf=FPDF()
    pdf.add_page()
    pdf.set_font("Arial",size=10)
    pdf.multi_cell(0,6,summary)
    for t in tasks:
        pdf.multi_cell(0,6,f"- {t.get('assignee')}: {t.get('task')} (do {t.get('deadline')})")
    pdf.output(path)
    return path
